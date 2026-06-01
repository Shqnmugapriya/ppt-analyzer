from pptx import Presentation
import os

def extract_text_from_ppt(file_path):
    """
    Extract titles and texts from PPT slides.
    """
    if not os.path.exists(file_path):
        print("ERROR: PPT file not found at:", os.path.abspath(file_path))
        return []

    prs = Presentation(file_path)
    ppt_data = []

    for i, slide in enumerate(prs.slides):
        slide_info = {"slide_number": i+1, "title": "", "texts": []}
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text.strip()
                if text:
                    if not slide_info["title"]:
                        slide_info["title"] = text
                    else:
                        slide_info["texts"].append(text)
        ppt_data.append(slide_info)
    
    return ppt_data